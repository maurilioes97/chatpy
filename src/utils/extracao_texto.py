from pathlib import Path


def extrair_texto_de_arquivo(caminho_arquivo):
    caminho = Path(caminho_arquivo)

    if caminho.suffix.lower() != ".xlsx":
        raise RuntimeError(
            "Formato de arquivo nao suportado. Envie apenas planilhas .xlsx."
        )

    return extrair_texto_xlsx(caminho)


def extrair_texto_xlsx(caminho_xlsx):
    try:
        from openpyxl import load_workbook
    except ImportError as erro:
        raise RuntimeError(
            "Para processar arquivos .xlsx, instale a dependencia openpyxl."
        ) from erro

    pasta_trabalho = load_workbook(filename=caminho_xlsx, data_only=True)
    blocos = []

    for aba in pasta_trabalho.worksheets:
        linhas = [f"Aba: {aba.title}"]

        for linha in aba.iter_rows(values_only=True):
            valores = []
            for celula in linha:
                if celula is None:
                    continue

                texto = str(celula).strip()
                if texto:
                    valores.append(texto)

            if valores:
                linhas.append(" | ".join(valores))

        if len(linhas) > 1:
            blocos.append("\n".join(linhas))

    return "\n\n".join(blocos).strip()
